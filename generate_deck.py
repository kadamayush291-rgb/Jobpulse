import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Clean Dark Palette
    COLOR_BG = RGBColor(15, 23, 42)          # Deep Dark Navy #0F172A
    COLOR_CARD_BG = RGBColor(30, 41, 59)     # Slate Navy Card #1E293B
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate Border #334155
    COLOR_CYAN = RGBColor(14, 165, 233)      # Accent Cyan #0EA5E9
    COLOR_PURPLE = RGBColor(129, 140, 248)   # Accent Purple #818CF8
    COLOR_AMBER = RGBColor(245, 158, 11)     # Accent Amber #F59E0B
    COLOR_GREEN = RGBColor(16, 185, 129)     # Accent Emerald #10B981
    COLOR_TEXT_WHITE = RGBColor(248, 250, 252) # Crisp White #F8FAFC
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # Muted Gray #94A3B8

    # Base directory for images
    img_dir = r"C:\Users\Ayush\.gemini\antigravity\brain\e3ed8295-9c80-401f-b121-678b7f210568"
    img_hierarchy = os.path.join(img_dir, "llm_hierarchy_1790606495558.jpg")
    img_transformer = os.path.join(img_dir, "transformer_architecture_1790606302463.jpg")
    img_training = os.path.join(img_dir, "llm_training_pipeline_1790606363219.jpg")
    img_next_token = os.path.join(img_dir, "next_token_prob_1790606411763.jpg")

    def set_slide_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, title_text, badge_text="PRESENTATION OVERVIEW"):
        # Category Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(12)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_CYAN
        p_b.font.name = "Arial"

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE
        p_t.font.name = "Arial"

    # ==========================================
    # SLIDE 1: Title Slide (Cover)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    add_card(slide1, Inches(0.8), Inches(1.2), Inches(11.733), Inches(2.6), bg_color=RGBColor(30, 41, 59), border_color=COLOR_CYAN)
    
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(11), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "COLLEGE PROJECT PRESENTATION"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN
    p0.space_after = Pt(8)
    
    p1 = tf.add_paragraph()
    p1.text = "How ChatGPT & Other LLMs Work"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "A Simple & Practical Overview of Large Language Models, Architecture & Training"
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_TEXT_MUTED

    # Student Info Box
    add_card(slide1, Inches(0.8), Inches(4.2), Inches(11.733), Inches(2.6), bg_color=RGBColor(24, 33, 47), border_color=COLOR_CARD_BORDER)
    
    infobox = slide1.shapes.add_textbox(Inches(1.2), Inches(4.4), Inches(11), Inches(2.2))
    itf = infobox.text_frame
    itf.word_wrap = True

    ip0 = itf.paragraphs[0]
    ip0.text = "PRESENTED BY"
    ip0.font.size = Pt(14)
    ip0.font.bold = True
    ip0.font.color.rgb = COLOR_PURPLE
    ip0.space_after = Pt(12)

    details = [
        ("Group Members / Authors:", " [Click to add Names - e.g., Student A, Student B]"),
        ("Roll Numbers / Section:", " [Click to add Roll Nos / Class - e.g., CS-2024-001]"),
        ("Course / Department:", " Computer Science & Engineering / Artificial Intelligence"),
        ("Project Guide / Instructor:", " [Click to add Guide Name]")
    ]

    for label, val in details:
        p = itf.add_paragraph()
        r1 = p.add_run()
        r1.text = label
        r1.font.bold = True
        r1.font.size = Pt(15)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(15)
        r2.font.color.rgb = COLOR_CYAN if "[" in val else COLOR_TEXT_MUTED
        p.space_after = Pt(5)

    # ==========================================
    # SLIDE 2: What is ChatGPT & LLMs? (Intro)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "What is ChatGPT & Large Language Models?", badge_text="01 | INTRODUCTION")

    # Image Left
    if os.path.exists(img_hierarchy):
        slide2.shapes.add_picture(img_hierarchy, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))

    # Right Card: Clear Breakdown
    add_card(slide2, Inches(6.7), Inches(1.8), Inches(5.833), Inches(5.0), border_color=COLOR_CYAN)
    rbox = slide2.shapes.add_textbox(Inches(6.9), Inches(2.0), Inches(5.433), Inches(4.6))
    rtf = rbox.text_frame
    rtf.word_wrap = True

    p0 = rtf.paragraphs[0]
    p0.text = "CORE DEFINITIONS & DIFFERENCE"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN
    p0.space_after = Pt(10)

    points = [
        ("LLM (Large Language Model): ", "An AI neural network trained on massive text data to predict language."),
        ("GPT (Generative Pre-trained Transformer): ", "OpenAI's specific model design built on Transformer layers."),
        ("ChatGPT: ", "The end-user app featuring a chat UI, conversation memory, and safety filters."),
        ("Rule-Based vs Generative: ", "Older bots used fixed IF-THEN scripts. ChatGPT generates brand new text word-by-word based on statistical patterns.")
    ]

    for title, desc in points:
        p = rtf.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: Core Concepts: Parameters & Tokens
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "What Makes an LLM 'Large'? Parameters & Tokens", badge_text="02 | CORE CONCEPTS")

    # Card 1: Parameters
    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    pbox = slide3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.6))
    ptf = pbox.text_frame
    ptf.word_wrap = True

    pp0 = ptf.paragraphs[0]
    pp0.text = "1. PARAMETERS (MODEL SIZE)"
    pp0.font.size = Pt(16)
    pp0.font.bold = True
    pp0.font.color.rgb = COLOR_CYAN
    pp0.space_after = Pt(12)

    p_items = [
        ("What is a Parameter? ", "An adjustable numerical weight inside the neural network (connection strength)."),
        ("The Knobs Analogy: ", "Think of parameters as tiny dials tuned millions of times during training."),
        ("Massive Scale: ", "Modern LLMs range from 7 Billion to 1.7+ Trillion parameters."),
        ("Impact: ", "More parameters = richer understanding of grammar, facts, and reasoning.")
    ]
    for title, desc in p_items:
        p = ptf.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # Card 2: Tokens
    add_card(slide3, Inches(6.833), Inches(1.8), Inches(5.7), Inches(5.0))
    tbox = slide3.shapes.add_textbox(Inches(7.133), Inches(2.0), Inches(5.1), Inches(4.6))
    ttf = tbox.text_frame
    ttf.word_wrap = True

    tp0 = ttf.paragraphs[0]
    tp0.text = "2. TOKENS (DATA SIZE)"
    tp0.font.size = Pt(16)
    tp0.font.bold = True
    tp0.font.color.rgb = COLOR_PURPLE
    tp0.space_after = Pt(12)

    t_items = [
        ("What is a Token? ", "Text is broken down into sub-word chunks (~4 characters in English)."),
        ("Easy Example: ", "'unbelievable' is split into 'un' + 'believ' + 'able' (3 tokens). Handles typos easily!"),
        ("Dataset Volume: ", "Trained on 10+ Trillion tokens from books, websites, and code."),
        ("Learning Goal: ", "Predict the hidden next token across trillions of sentences.")
    ]
    for title, desc in t_items:
        p = ttf.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 4: The Engine: Transformer & Self-Attention
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Inside the Transformer Engine & Self-Attention", badge_text="03 | ARCHITECTURE")

    # Image Left
    if os.path.exists(img_transformer):
        slide4.shapes.add_picture(img_transformer, Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))

    # Right Card: Key Highlights
    add_card(slide4, Inches(6.9), Inches(1.8), Inches(5.633), Inches(5.0), border_color=COLOR_CYAN)
    rbox4 = slide4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.233), Inches(4.6))
    rtf4 = rbox4.text_frame
    rtf4.word_wrap = True

    p0 = rtf4.paragraphs[0]
    p0.text = "KEY ARCHITECTURAL HIGHLIGHTS"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN
    p0.space_after = Pt(10)

    arch_points = [
        ("Token Embeddings: ", "Converts token IDs into mathematical vector representations."),
        ("Positional Encoding: ", "Adds order info so model distinguishes 'Dog bites man' vs 'Man bites dog'."),
        ("Self-Attention Mechanism: ", "Calculates focus scores between all words simultaneously."),
        ("Context Example: ", "In 'The trophy didn't fit in the suitcase because it was too big', Self-Attention links 'it' -> 'trophy'.")
    ]
    for title, desc in arch_points:
        p = rtf4.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: How LLMs Are Trained (3 Stages)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "How LLMs Are Trained: The 3 Stages", badge_text="04 | TRAINING PIPELINE")

    # Image Top
    if os.path.exists(img_training):
        slide5.shapes.add_picture(img_training, Inches(0.8), Inches(1.7), Inches(11.733), Inches(2.6))

    # 3 Stage Cards Below
    stages = [
        ("STAGE 1: PRE-TRAINING", "Base Model Creation", "Trained on trillions of text tokens. Learns grammar, facts & logic by predicting next tokens.", COLOR_CYAN),
        ("STAGE 2: FINE-TUNING (SFT)", "Instruction Following", "Trained on curated human Q&A examples. Teaches raw text model to behave like a helpful assistant.", COLOR_PURPLE),
        ("STAGE 3: RLHF ALIGNMENT", "Human Safety & Feedback", "Human reviewers rank answers. A reward model guides training for safety, clarity & honesty.", COLOR_GREEN)
    ]

    for i, (stitle, ssub, sdesc, saccent) in enumerate(stages):
        left = Inches(0.8 + i * 3.98)
        add_card(slide5, left, Inches(4.5), Inches(3.78), Inches(2.4), border_color=saccent)
        
        sbox = slide5.shapes.add_textbox(left + Inches(0.15), Inches(4.65), Inches(3.48), Inches(2.1))
        stf = sbox.text_frame
        stf.word_wrap = True
        
        sp0 = stf.paragraphs[0]
        sp0.text = stitle
        sp0.font.size = Pt(12)
        sp0.font.bold = True
        sp0.font.color.rgb = saccent
        sp0.space_after = Pt(2)

        sp1 = stf.add_paragraph()
        sp1.text = ssub
        sp1.font.size = Pt(15)
        sp1.font.bold = True
        sp1.font.color.rgb = COLOR_TEXT_WHITE
        sp1.space_after = Pt(6)

        sp2 = stf.add_paragraph()
        sp2.text = sdesc
        sp2.font.size = Pt(12)
        sp2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 6: How Responses Are Generated (Inference)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "How ChatGPT Generates Answers (Inference)", badge_text="05 | RESPONSE GENERATION")

    # Left Card: Process
    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    ibox = slide6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    itf = ibox.text_frame
    itf.word_wrap = True

    ip0 = itf.paragraphs[0]
    ip0.text = "AUTOREGRESSIVE GENERATION"
    ip0.font.size = Pt(16)
    ip0.font.bold = True
    ip0.font.color.rgb = COLOR_CYAN
    ip0.space_after = Pt(10)

    i_steps = [
        ("1. Read Context: ", "Reads prompt + past chat history inside the context window."),
        ("2. Calculate Odds: ", "Outputs probabilities for next token. E.g. 'The cat sat on the ___' -> mat (62%), floor (21%), chair (9%)."),
        ("3. Temperature Setting: ", "Low temperature = factual & predictable; High temperature = creative & adventurous."),
        ("4. Repeat Word-by-Word: ", "Selected token is added to sequence and steps repeat until finished.")
    ]
    for t, d in i_steps:
        p = itf.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + t
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # Image Right
    if os.path.exists(img_next_token):
        slide6.shapes.add_picture(img_next_token, Inches(6.833), Inches(1.8), Inches(5.7), Inches(5.0))

    # ==========================================
    # SLIDE 7: Models, Uses & Critical Limitations
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Popular LLMs, Applications & Key Limitations", badge_text="06 | LANDSCAPE & LIMITS")

    # Left Card: Popular Models & Uses
    add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0), border_color=COLOR_CYAN)
    lbox7 = slide7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    ltf7 = lbox7.text_frame
    ltf7.word_wrap = True

    lp0 = ltf7.paragraphs[0]
    lp0.text = "POPULAR MODELS & APPLICATIONS"
    lp0.font.size = Pt(16)
    lp0.font.bold = True
    lp0.font.color.rgb = COLOR_CYAN
    lp0.space_after = Pt(10)

    models_uses = [
        ("Leading Models: ", "OpenAI GPT-4o, Anthropic Claude 3.5, Google Gemini 1.5, Meta Llama 3."),
        ("Coding & Tech: ", "Writing functions, debugging errors, explaining codebases, and writing tests."),
        ("Writing & Editing: ", "Drafting essays, polishing tone, summarizing long papers."),
        ("Customer Support: ", "Powering 24/7 smart bots with multi-turn conversation memory.")
    ]
    for m, d in models_uses:
        p = ltf7.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + m
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # Right Card: Key Limitations
    add_card(slide7, Inches(6.833), Inches(1.8), Inches(5.7), Inches(5.0), border_color=COLOR_AMBER)
    rbox7 = slide7.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.6))
    rtf7 = rbox7.text_frame
    rtf7.word_wrap = True

    rp0 = rtf7.paragraphs[0]
    rp0.text = "CRITICAL LIMITATIONS"
    rp0.font.size = Pt(16)
    rp0.font.bold = True
    rp0.font.color.rgb = COLOR_AMBER
    rp0.space_after = Pt(10)

    limits = [
        ("Hallucinations: ", "Models predict word odds—not verified facts. They can state false facts confidently. (Mitigation: Web Search/RAG)."),
        ("Knowledge Cutoff: ", "Unaware of events after training dataset date limit."),
        ("Human Biases: ", "Learns imbalances present in internet text datasets."),
        ("No True Understanding: ", "Statistical pattern matching != human consciousness or real-world reasoning.")
    ]
    for l, d in limits:
        p = rtf7.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + l
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: Summary & Viva Q&A Guide
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Key Takeaways & Viva Q&A Quick Reference", badge_text="07 | CONCLUSION & VIVA PREP")

    # Left Card: Summary
    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0), border_color=COLOR_CYAN)
    sbox8 = slide8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    stf8 = sbox8.text_frame
    stf8.word_wrap = True

    sp0 = stf8.paragraphs[0]
    sp0.text = "3 CORE TAKEAWAYS"
    sp0.font.size = Pt(16)
    sp0.font.bold = True
    sp0.font.color.rgb = COLOR_CYAN
    sp0.space_after = Pt(10)

    sum_items = [
        ("Statistical Word Predictor: ", "LLMs generate probable next tokens based on patterns learned from trillions of sentences."),
        ("Transformer Breakthrough: ", "Self-Attention processes full sentences in parallel on GPUs to understand word context."),
        ("RLHF Alignment: ", "Pre-training learns raw text; Fine-Tuning and RLHF turn it into a safe assistant.")
    ]
    for t, d in sum_items:
        p = stf8.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + t
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(14)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # Right Card: Viva Q&A Cheat Sheet
    add_card(slide8, Inches(6.833), Inches(1.8), Inches(5.7), Inches(5.0), border_color=COLOR_AMBER)
    vbox8 = slide8.shapes.add_textbox(Inches(7.033), Inches(2.0), Inches(5.3), Inches(4.6))
    vtf8 = vbox8.text_frame
    vtf8.word_wrap = True

    vp0 = vtf8.paragraphs[0]
    vp0.text = "TOP VIVA / Q&A QUESTIONS"
    vp0.font.size = Pt(16)
    vp0.font.bold = True
    vp0.font.color.rgb = COLOR_AMBER
    vp0.space_after = Pt(10)

    viva_qna = [
        ("Q1: Pre-training vs. Fine-Tuning?", "Pre-training teaches general language knowledge; Fine-Tuning teaches instruction-following behavior."),
        ("Q2: What is Temperature?", "Controls response randomness: lower values are factual/precise, higher values are creative/varied.")
    ]
    for q, a in viva_qna:
        p = vtf8.add_paragraph()
        r1 = p.add_run()
        r1.text = q + "\n"
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_WHITE
        r2 = p.add_run()
        r2.text = "Ans: " + a
        r2.font.size = Pt(13.5)
        r2.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(8)

    # Output file paths (save to main + simple fallback)
    output_path_simple = r"c:\Users\Ayush\OneDrive\Desktop\DE project\How_ChatGPT_and_LLMs_Work_Simple.pptx"
    prs.save(output_path_simple)
    print(f"Saved to: {output_path_simple}")

    output_path_main = r"c:\Users\Ayush\OneDrive\Desktop\DE project\How_ChatGPT_and_LLMs_Work.pptx"
    try:
        prs.save(output_path_main)
        print(f"Saved to main: {output_path_main}")
    except Exception as e:
        print(f"Could not overwrite main file because it is open, but simple file saved: {e}")

    # Also save copy in brain artifact directory
    artifact_path = os.path.join(img_dir, "How_ChatGPT_and_LLMs_Work_Simple.pptx")
    prs.save(artifact_path)
    print(f"Artifact presentation saved to: {artifact_path}")

if __name__ == "__main__":
    create_deck()
